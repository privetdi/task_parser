from datetime import date, datetime
from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

from app.services.report import HEADER_ROW, ReportRow, build_workbook

REPORT_DATE = date(2026, 10, 22)
GENERATED = datetime(2026, 10, 22, 9, 0, tzinfo=ZoneInfo("Europe/Minsk"))

ROWS = [
    ReportRow("AZN", "Азербайджанский манат", 1, 1.7769, date(2026, 10, 1), 1, -0.011),
    ReportRow("RUB", "Российский рубль", 100, 3.5, date(2026, 10, 22), 0, 0.0),
    ReportRow("USD", "Доллар США", 1, 2.9713, date(2026, 10, 22), 0, None),
]


def open_report(rows: list[ReportRow] = ROWS):
    content = build_workbook(rows, REPORT_DATE, GENERATED)
    return load_workbook(BytesIO(content)).active


def test_title_shows_report_date_and_currency_count() -> None:
    sheet = open_report()

    assert "22.10.2026" in sheet["A1"].value
    assert sheet["A2"].value == "Валют в отчёте: 3"
    assert "09:00" in sheet["A3"].value and "Europe/Minsk" in sheet["A3"].value
    assert "nbrb.by" in sheet["A4"].value


def test_table_has_header_and_one_row_per_currency() -> None:
    sheet = open_report()

    headers = [c.value for c in sheet[HEADER_ROW]]
    assert headers == [
        "Код", "Валюта", "Номинал", "Курс, BYN", "Дата курса", "Периодичность", "Изменение",
    ]
    data = [[c.value for c in sheet[HEADER_ROW + i]] for i in range(1, len(ROWS) + 1)]
    assert [r[0] for r in data] == ["AZN", "RUB", "USD"]
    assert sheet.max_row == HEADER_ROW + len(ROWS)
    # Число строк таблицы совпадает с числом из заголовка
    assert sheet.max_row - HEADER_ROW == int(sheet["A2"].value.rsplit(" ", 1)[1])


def test_every_row_has_its_own_rate_date_as_a_real_date() -> None:
    sheet = open_report()

    dates = [sheet.cell(row=HEADER_ROW + i, column=5).value for i in range(1, len(ROWS) + 1)]
    assert [d.date() for d in dates] == [date(2026, 10, 1), date(2026, 10, 22), date(2026, 10, 22)]
    assert sheet.cell(row=HEADER_ROW + 1, column=5).number_format == "DD.MM.YYYY"


def test_values_scale_periodicity_and_change_are_written() -> None:
    sheet = open_report()
    azn, rub, usd = ([c.value for c in sheet[HEADER_ROW + i]] for i in (1, 2, 3))

    assert azn[3] == 1.7769 and azn[5] == "ежемесячно" and azn[6] == -0.011
    assert rub[2] == 100 and rub[5] == "ежедневно"
    assert usd[6] is None  # предыдущего значения нет


def test_header_is_frozen_and_table_is_filterable() -> None:
    sheet = open_report()

    assert sheet.freeze_panes == f"A{HEADER_ROW + 1}"
    assert sheet.auto_filter.ref == f"A{HEADER_ROW}:G{HEADER_ROW + len(ROWS)}"


def test_empty_report_is_still_a_valid_workbook() -> None:
    sheet = open_report([])

    assert sheet["A2"].value == "Валют в отчёте: 0"
    assert sheet.max_row == HEADER_ROW


def test_print_layout_fits_all_columns_on_the_page_and_repeats_the_header() -> None:
    sheet = open_report()

    assert sheet.page_setup.orientation == "landscape"
    assert sheet.sheet_properties.pageSetUpPr.fitToPage is True
    assert sheet.page_setup.fitToWidth == 1 and sheet.page_setup.fitToHeight == 0
    # openpyxl делает ссылку абсолютной
    assert sheet.print_title_rows == f"${HEADER_ROW}:${HEADER_ROW}"
