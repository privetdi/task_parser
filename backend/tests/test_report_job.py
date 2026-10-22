from datetime import UTC, date, datetime
from io import BytesIO
from zoneinfo import ZoneInfo

import pytest
from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.db import init_db, make_engine, make_session_factory
from app.models import Currency, Delivery, DeliveryKind, DeliveryStatus, ExchangeRate, JobStatus
from app.services.deliveries import get_delivery
from app.services.mailer import MailError
from app.services.report import HEADER_ROW, render_report, send_rates_report
from tests.helpers import RecordingMailer

DAY = date(2026, 10, 22)
YESTERDAY = date(2026, 10, 21)
GENERATED = datetime(2026, 10, 22, 9, 0, tzinfo=ZoneInfo("Europe/Minsk"))
SENT_AT = datetime(2026, 10, 22, 6, 0, tzinfo=UTC)


@pytest.fixture(autouse=True)
def rates(session: Session) -> None:
    for cur_id, parent, code in [(431, 145, "USD"), (451, 19, "EUR")]:
        session.add(Currency(
            cur_id=cur_id, parent_id=parent, abbreviation=code, name=code, scale=1,
            periodicity=0, date_start=date(2021, 7, 9), date_end=date(2050, 1, 1),
        ))
        session.add(ExchangeRate(currency_id=cur_id, rate_date=YESTERDAY, official_rate=2.9))
        session.add(ExchangeRate(currency_id=cur_id, rate_date=DAY, official_rate=3.0))
    session.commit()


@pytest.fixture
def mailer() -> RecordingMailer:
    return RecordingMailer()


def send(session: Session, mailer: RecordingMailer, *, force: bool = False, day: date = DAY):
    return send_rates_report(
        session, mailer, report_date=day, generated_at=GENERATED, force=force, now=lambda: SENT_AT
    )


def delivery_count(session: Session) -> int:
    return session.scalar(select(func.count()).select_from(Delivery))


def test_scheduled_send_delivers_readable_xlsx_with_date_and_currencies(
    session: Session, mailer: RecordingMailer
) -> None:
    result = send(session, mailer)

    assert result.status == JobStatus.SUCCESS and "валют 2" in result.summary
    (mail,) = mailer.sent
    assert mail.subject == "Курсы валют НБ РБ на 22.10.2026"
    assert "EUR, USD" in mail.body and "Валют в отчёте: 2" in mail.body
    assert "на 22.10.2026" in mail.body
    (attachment,) = mail.attachments
    assert attachment.filename == "rates_2026-10-22.xlsx"
    sheet = load_workbook(BytesIO(attachment.content)).active
    assert "22.10.2026" in sheet["A1"].value
    assert sheet.cell(row=HEADER_ROW + 1, column=1).value == "EUR"
    delivery = get_delivery(session, DeliveryKind.RATES_REPORT, "2026-10-22")
    assert delivery.status == DeliveryStatus.SENT and delivery.sent_at == SENT_AT


def test_second_scheduled_run_for_same_date_is_skipped(
    session: Session, mailer: RecordingMailer
) -> None:
    send(session, mailer)

    result = send(session, mailer)

    assert result.status == JobStatus.SKIPPED
    assert "уже отправлен" in result.summary
    assert len(mailer.sent) == 1 and delivery_count(session) == 1


def test_manual_send_goes_out_even_if_already_sent_today(
    session: Session, mailer: RecordingMailer
) -> None:
    send(session, mailer)

    result = send(session, mailer, force=True)

    assert result.status == JobStatus.SUCCESS
    assert len(mailer.sent) == 2 and delivery_count(session) == 1


def test_mail_failure_does_not_mark_report_as_sent_and_next_run_retries(
    session: Session, mailer: RecordingMailer
) -> None:
    mailer.error = MailError("SMTP-сервер недоступен")

    failed = send(session, mailer)

    assert failed.status == JobStatus.ERROR and "недоступен" in failed.errors[0]
    delivery = get_delivery(session, DeliveryKind.RATES_REPORT, "2026-10-22")
    assert delivery.status == DeliveryStatus.FAILED and delivery.sent_at is None

    mailer.error = None
    retried = send(session, mailer)

    assert retried.status == JobStatus.SUCCESS and len(mailer.sent) == 1


def test_next_day_report_is_sent_again(session: Session, mailer: RecordingMailer) -> None:
    send(session, mailer)

    result = send(session, mailer, day=date(2026, 10, 23))

    assert result.status == JobStatus.SUCCESS
    assert len(mailer.sent) == 2 and delivery_count(session) == 2


def test_empty_database_sends_nothing_and_reports_why(mailer: RecordingMailer) -> None:
    empty = make_engine("sqlite://")
    init_db(empty)
    with make_session_factory(empty)() as empty_session:
        result = send(empty_session, mailer)
        assert delivery_count(empty_session) == 0

    assert result.status == JobStatus.ERROR
    assert "нет курсов" in result.summary and "загрузку" in result.errors[0]
    assert mailer.sent == []


def test_render_report_gives_the_same_file_without_sending_anything(
    session: Session, mailer: RecordingMailer
) -> None:
    report = render_report(session, DAY, GENERATED)

    assert report.filename == "rates_2026-10-22.xlsx"
    assert [row.code for row in report.rows] == ["EUR", "USD"]
    assert load_workbook(BytesIO(report.content)).active["A2"].value == "Валют в отчёте: 2"
    assert delivery_count(session) == 0 and mailer.sent == []
