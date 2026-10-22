"""Excel-отчёт с актуальными курсами НБ РБ."""
from collections.abc import Callable
from dataclasses import dataclass
from datetime import date, datetime
from io import BytesIO
from itertools import groupby

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.models import Currency, DeliveryKind, ExchangeRate, JobStatus, utcnow
from app.services.deliveries import Outcome, deliver_force, deliver_once, get_delivery
from app.services.mailer import XLSX_MIME, Attachment, Mailer, OutgoingMail
from app.services.result import JobResult

SOURCE_URL = "https://www.nbrb.by/statistics/rates/ratesdaily"
PERIODICITY_LABELS = {0: "ежедневно", 1: "ежемесячно"}
# Колонки таблицы: (заголовок, ширина, формат числа)
COLUMNS = [
    ("Код", 8, None),
    ("Валюта", 40, None),
    ("Номинал", 10, "0"),
    ("Курс, BYN", 14, "0.0000"),
    ("Дата курса", 14, "DD.MM.YYYY"),
    ("Периодичность", 16, None),
    ("Изменение", 12, "+0.0000;-0.0000;0.0000"),
]
HEADER_ROW = 7  # выше — заголовок отчёта и пояснения


@dataclass(frozen=True)
class ReportRow:
    code: str
    name: str
    scale: int  # курс указан за столько единиц валюты
    rate: float  # официальный курс, BYN
    rate_date: date  # дата, к которой относится курс
    periodicity: int  # 0 — ежедневно, 1 — ежемесячно
    change: float | None  # изменение к предыдущему значению (None, если предыдущего нет)


def collect_report_rows(session: Session, report_date: date) -> list[ReportRow]:
    """Для каждой действующей валюты — самый свежий курс с датой не позже `report_date`.

    Версии одной валюты (общий parent_id) склеиваются. Курсы с датой позже даты отчёта
    (например, уже опубликованный курс «на завтра») не учитываются. Валюты, по которым
    в базе нет ни одного курса, в отчёт не попадают.
    """
    active_parents = select(Currency.parent_id).where(Currency.date_end >= report_date)
    ranked = (
        select(
            Currency.parent_id.label("parent_id"),
            Currency.abbreviation.label("code"),
            Currency.name.label("name"),
            Currency.scale.label("scale"),
            Currency.periodicity.label("periodicity"),
            ExchangeRate.rate_date.label("rate_date"),
            ExchangeRate.official_rate.label("rate"),
            func.row_number()
            .over(partition_by=Currency.parent_id, order_by=ExchangeRate.rate_date.desc())
            .label("rn"),
        )
        .join(ExchangeRate, ExchangeRate.currency_id == Currency.cur_id)
        .where(ExchangeRate.rate_date <= report_date, Currency.parent_id.in_(active_parents))
        .subquery()
    )
    # rn=1 — последний курс, rn=2 — предыдущий (нужен для «изменения»)
    records = session.execute(
        select(ranked).where(ranked.c.rn <= 2).order_by(ranked.c.parent_id, ranked.c.rn)
    ).all()

    rows = []
    for _, group in groupby(records, key=lambda record: record.parent_id):
        latest, *previous = list(group)
        change = None
        if previous:
            # У версий валюты номинал может различаться — сравниваем курс за одну единицу
            previous_in_latest_scale = previous[0].rate / previous[0].scale * latest.scale
            # Округляем, чтобы в ячейке не оказалось float-шума вида -0.04630000000000045
            change = round(latest.rate - previous_in_latest_scale, 6)
        rows.append(
            ReportRow(
                code=latest.code,
                name=latest.name,
                scale=latest.scale,
                rate=latest.rate,
                rate_date=latest.rate_date,
                periodicity=latest.periodicity,
                change=change,
            )
        )
    return sorted(rows, key=lambda row: row.code)


def build_workbook(rows: list[ReportRow], report_date: date, generated_at: datetime) -> bytes:
    """Собирает .xlsx: заголовок с датой и числом валют, под ним таблица курсов.

    Дата курса указана в каждой строке отдельно: у ежемесячных валют она отличается от
    даты отчёта, и по ней видно, к какому дню относятся данные.
    """
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Курсы"

    timezone_name = getattr(generated_at.tzinfo, "key", None) or generated_at.tzname() or ""
    last_column = get_column_letter(len(COLUMNS))

    sheet["A1"] = f"Официальные курсы валют НБ РБ на {report_date:%d.%m.%Y}"
    sheet["A1"].font = Font(bold=True, size=14)
    sheet.merge_cells(f"A1:{last_column}1")
    sheet["A2"] = f"Валют в отчёте: {len(rows)}"
    generated = f"Сформировано: {generated_at:%d.%m.%Y %H:%M}"
    sheet["A3"] = f"{generated} ({timezone_name})" if timezone_name else generated
    sheet["A4"] = f"Источник: {SOURCE_URL}"
    sheet["A4"].hyperlink = SOURCE_URL
    sheet["A5"] = (
        "Курс указан в белорусских рублях (BYN) за указанный номинал. «Дата курса» — день, "
        "к которому он относится; у ежемесячных валют это дата последнего установленного курса."
    )
    sheet["A5"].alignment = Alignment(wrap_text=True, vertical="top")
    sheet.merge_cells(f"A5:{last_column}5")
    sheet.row_dimensions[5].height = 32

    border = Border(*(Side(style="thin", color="BBBBBB"),) * 4)
    for column, (title, width, _) in enumerate(COLUMNS, start=1):
        cell = sheet.cell(row=HEADER_ROW, column=column, value=title)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DCE6F1")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
        sheet.column_dimensions[get_column_letter(column)].width = width

    for offset, row in enumerate(rows, start=1):
        values = [
            row.code,
            row.name,
            row.scale,
            row.rate,
            row.rate_date,
            PERIODICITY_LABELS.get(row.periodicity, str(row.periodicity)),
            row.change,
        ]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=HEADER_ROW + offset, column=column, value=value)
            cell.border = border
            number_format = COLUMNS[column - 1][2]
            if number_format:
                cell.number_format = number_format

    sheet.freeze_panes = sheet.cell(row=HEADER_ROW + 1, column=1)
    # Печать: все колонки по ширине страницы, шапка таблицы повторяется на каждой странице
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    sheet.print_title_rows = f"{HEADER_ROW}:{HEADER_ROW}"
    sheet.auto_filter.ref = f"A{HEADER_ROW}:{last_column}{HEADER_ROW + max(len(rows), 1)}"

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


@dataclass(frozen=True)
class RenderedReport:
    content: bytes
    filename: str
    rows: list[ReportRow]


def render_report(session: Session, report_date: date, generated_at: datetime) -> RenderedReport:
    """Собирает отчёт без отправки (его же отдаёт кнопка «Скачать Excel»)."""
    rows = collect_report_rows(session, report_date)
    return RenderedReport(
        content=build_workbook(rows, report_date, generated_at),
        filename=f"rates_{report_date.isoformat()}.xlsx",
        rows=rows,
    )


def _compose_mail(
    report: RenderedReport, report_date: date, generated_at: datetime
) -> OutgoingMail:
    newest = max(row.rate_date for row in report.rows)
    body = (
        f"Официальные курсы валют НБ РБ на {report_date:%d.%m.%Y}.\n"
        f"Валют в отчёте: {len(report.rows)}. Самый свежий курс — на {newest:%d.%m.%Y}; "
        f"у каждой валюты в таблице указана дата, к которой относится её курс.\n\n"
        f"Валюты: {', '.join(row.code for row in report.rows)}.\n\n"
        f"Сформировано: {generated_at:%d.%m.%Y %H:%M}. Файл — во вложении."
    )
    return OutgoingMail(
        subject=f"Курсы валют НБ РБ на {report_date:%d.%m.%Y}",
        body=body,
        attachments=(Attachment(report.filename, report.content, XLSX_MIME),),
    )


def send_rates_report(
    session: Session,
    mailer: Mailer,
    *,
    report_date: date,
    generated_at: datetime,
    force: bool,
    now: Callable[[], datetime] = utcnow,
) -> JobResult:
    """Формирует отчёт и отправляет его.

    По расписанию (`force=False`) отчёт за дату уходит не больше одного раза, в том числе
    после перезапуска сервиса. Ручной запуск (`force=True`) отправляет всегда.
    """
    report = render_report(session, report_date, generated_at)
    if not report.rows:
        return JobResult(
            JobStatus.ERROR,
            "Отчёт не отправлен: в базе нет курсов",
            ["Нет данных для отчёта — сначала выполните загрузку курсов"],
        )

    deliver = deliver_force if force else deliver_once
    result = deliver(
        session,
        mailer,
        DeliveryKind.RATES_REPORT,
        report_date.isoformat(),
        lambda: _compose_mail(report, report_date, generated_at),
        now,
    )
    if result.outcome == Outcome.SKIPPED:
        sent_at = get_delivery(session, DeliveryKind.RATES_REPORT, report_date.isoformat()).sent_at
        return JobResult(
            JobStatus.SKIPPED,
            f"Отчёт за {report_date:%d.%m.%Y} уже отправлен ({sent_at:%d.%m.%Y %H:%M} UTC)",
        )
    if result.outcome == Outcome.FAILED:
        return JobResult(
            JobStatus.ERROR,
            f"Отчёт за {report_date:%d.%m.%Y} не отправлен",
            [result.detail],
        )
    return JobResult(
        JobStatus.SUCCESS,
        f"Отчёт за {report_date:%d.%m.%Y} отправлен: валют {len(report.rows)}; {result.detail}",
    )
